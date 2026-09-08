#!/usr/bin/env python3
import argparse, json
from pathlib import Path
from shapely.geometry import shape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def route_num(rid):
    try: return int(str(rid).lstrip('Rr'))
    except: return 10**9

def main():
    ap=argparse.ArgumentParser(description='Assign final PortWatch chokepoints to 1,037 Eurostat SeaRoute geometries by direct geometric intersection.')
    ap.add_argument('routes_geojson')
    ap.add_argument('chokepoints_geojson')
    ap.add_argument('--out-prefix',default='LNG_1037_routes_with_final_chokepoints')
    args=ap.parse_args()
    with open(args.routes_geojson,encoding='utf-8-sig') as f: routes=json.load(f)
    with open(args.chokepoints_geojson,encoding='utf-8-sig') as f: cpj=json.load(f)
    cps=[(ft['properties'].get('node_id'),ft['properties'].get('chokepoint'),shape(ft['geometry'])) for ft in cpj['features']]
    if len(cps)!=28: raise ValueError(f'Expected 28 chokepoints, got {len(cps)}')
    feats=routes.get('features',[])
    if len(feats)!=1037: raise ValueError(f'Expected 1,037 route features, got {len(feats)}')
    seen=set(); rc=[]; qa=[]
    for ft in feats:
        p=ft.setdefault('properties',{})
        rid=p.get('route_id')
        if not rid: raise ValueError('A route feature has no route_id')
        if rid in seen: raise ValueError(f'Duplicate route_id {rid}')
        seen.add(rid)
        g=shape(ft['geometry'])
        hits=[]
        for node_id,name,cpg in cps:
            if g.intersects(cpg):
                hits.append((node_id,name))
                rc.append({'route_id':rid,'chokepoint_node_id':node_id,'chokepoint':name})
        p['n_chokepoints_final']=len(hits)
        p['chokepoints_final']='; '.join(name for _,name in hits)
        dist=p.get('distKM'); obs=p.get('voyage_distance_km_observed')
        try:
            diff=float(dist)-float(obs); ape=abs(diff)/float(obs) if float(obs)!=0 else None
        except: diff=ape=None
        try: dfrom=float(p.get('dFromKM'))
        except: dfrom=None
        try: dto=float(p.get('dToKM'))
        except: dto=None
        flag=[]
        if dfrom is not None and dfrom>100: flag.append('large_origin_snap')
        if dto is not None and dto>100: flag.append('large_destination_snap')
        if ape is not None and ape>0.15: flag.append('distance_deviation_gt15pct')
        qa.append({'route_id':rid,'distKM':dist,'observed_km':obs,'difference_km':diff,'abs_pct_error':ape,
                   'dFromKM':dfrom,'dToKM':dto,'n_chokepoints_final':len(hits),'qa_flags':'; '.join(flag)})
    feats.sort(key=lambda ft: route_num(ft.get('properties',{}).get('route_id')))
    out_geo=Path(args.out_prefix+'.geojson')
    with open(out_geo,'w',encoding='utf-8') as f: json.dump(routes,f,ensure_ascii=False,separators=(',',':'))
    wb=Workbook(); ws=wb.active; ws.title='Routes'
    prop_keys=[]
    for ft in feats:
        for k in ft.get('properties',{}):
            if k not in prop_keys: prop_keys.append(k)
    ws.append(prop_keys)
    for ft in feats: ws.append([ft.get('properties',{}).get(k) for k in prop_keys])
    wr=wb.create_sheet('Route-Chokepoint'); wr.append(['route_id','chokepoint_node_id','chokepoint'])
    for r in rc: wr.append(list(r.values()))
    wq=wb.create_sheet('QA'); qh=list(qa[0].keys()); wq.append(qh)
    for r in sorted(qa,key=lambda x:route_num(x['route_id'])): wq.append([r[k] for k in qh])
    wm=wb.create_sheet('Method');
    for row in [
        ('Item','Description'),
        ('Transit classification','Direct topological intersection between each Eurostat SeaRoute geometry and each of the 28 final chokepoint Polygon/MultiPolygon geometries.'),
        ('Distance buffers','None used for final transit classification.'),
        ('QA','Flags origin/destination network snap >100 km and route-vs-observed distance deviation >15% for manual review; flags do not automatically alter classifications.'),
    ]: wm.append(row)
    fill=PatternFill('solid',fgColor='1F4E78'); font=Font(color='FFFFFF',bold=True)
    for sh in wb.worksheets:
        sh.freeze_panes='A2'
        for c in sh[1]: c.fill=fill; c.font=font; c.alignment=Alignment(horizontal='center')
        for col in range(1,sh.max_column+1):
            vals=[str(sh.cell(r,col).value or '') for r in range(1,min(sh.max_row,200)+1)]
            sh.column_dimensions[get_column_letter(col)].width=min(max(10,max(map(len,vals))+2),45)
        for row in sh.iter_rows():
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    out_xlsx=Path(args.out_prefix+'.xlsx'); wb.save(out_xlsx)
    print(f'Wrote {out_geo} and {out_xlsx}; route-chokepoint rows={len(rc)}; flagged routes={sum(bool(x["qa_flags"]) for x in qa)}')
if __name__=='__main__': main()
