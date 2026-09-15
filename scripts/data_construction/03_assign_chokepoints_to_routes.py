#!/usr/bin/env python3
import argparse, json
from pathlib import Path
from pyproj import CRS, Transformer
from shapely.geometry import shape
from shapely.ops import transform
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BUFFER_KM = 25

def metric_buffer(geometry, distance_km):
    """Buffer a WGS84 geometry in metres using a local AEQD projection."""
    centre = geometry.representative_point()
    local = CRS.from_proj4(
        f"+proj=aeqd +lat_0={centre.y} +lon_0={centre.x} "
        "+datum=WGS84 +units=m +no_defs"
    )
    forward = Transformer.from_crs("OGC:CRS84", local, always_xy=True).transform
    inverse = Transformer.from_crs(local, "OGC:CRS84", always_xy=True).transform
    return transform(inverse, transform(forward, geometry).buffer(distance_km * 1000))

def route_num(rid):
    try: return int(str(rid).lstrip('Rr'))
    except: return 10**9

def main():
    
    project_root = Path(__file__).resolve().parents[2]
    data_dir = project_root / "data"

    ap = argparse.ArgumentParser(
        description=(
            "Assign final PortWatch chokepoints to 1,037 Eurostat "
            "SeaRoute geometries using a uniform 25 km spatial-tolerance buffer."
        )
    )

    ap.add_argument(
        "--routes",
        type=Path,
        default=data_dir / "LNG_1037_routes_searoute.geojson",
    )
    ap.add_argument(
        "--chokepoints",
        type=Path,
        default=data_dir / "PortWatch_28_chokepoints_geometry.geojson",
    )
    ap.add_argument(
        "--out-prefix",
        type=Path,
        default=data_dir / "LNG_1037_routes_with_final_chokepoints",
    )

    args = ap.parse_args()
    data_dir.mkdir(parents=True, exist_ok=True)

    with args.routes.open(encoding="utf-8-sig") as f:
        routes = json.load(f)

    with args.chokepoints.open(encoding="utf-8-sig") as f:
        cpj = json.load(f)
        
    cps=[]
    for ft in cpj['features']:
        geometry=shape(ft['geometry'])
        cps.append((
            ft['properties'].get('node_id'),
            ft['properties'].get('chokepoint'),
            geometry,
            metric_buffer(geometry, BUFFER_KM),
        ))
    if len(cps)!=28: raise ValueError(f'Expected 28 chokepoints, got {len(cps)}')
    
    feats=routes.get('features',[])
    if len(feats)!=1037: raise ValueError(f'Expected 1,037 route features, got {len(feats)}')
    seen=set(); rc=[]; qa=[]
    for ft in feats:
        p=ft.setdefault('properties',{})
        rid=p.get('route_id')
        if not rid: raise ValueError('A route feature has no route_id')
        if rid in seen: raise ValueError(f'Duplicate route_id {rid}')
        seen.add(rid)
        g=shape(ft['geometry'])
        exact_hits=[]; hits=[]
        for node_id,name,cpg,cpg_buffered in cps:
            exact=g.intersects(cpg)
            buffered=g.intersects(cpg_buffered)
            if exact: exact_hits.append((node_id,name))
            if buffered:
                hits.append((node_id,name))
                rc.append({
                    'route_id':rid,
                    'chokepoint_node_id':node_id,
                    'chokepoint':name,
                    'exact_intersection':exact,
                    'classification_method':'exact_intersection' if exact else f'within_{BUFFER_KM}km',
                })
        p['n_chokepoints_exact']=len(exact_hits)
        p['chokepoints_exact']='; '.join(name for _,name in exact_hits)
        p['n_chokepoints_final']=len(hits)
        p['chokepoints_final']='; '.join(name for _,name in hits)
        p['chokepoints_buffer_km']=BUFFER_KM
        p['n_chokepoints_added_by_buffer']=len(hits)-len(exact_hits)
        dist=p.get('distKM'); obs=p.get('voyage_distance_km_observed')
        try:
            diff=float(dist)-float(obs); ape=abs(diff)/float(obs) if float(obs)!=0 else None
        except: diff=ape=None
        try: dfrom=float(p.get('dFromKM'))
        except: dfrom=None
        try: dto=float(p.get('dToKM'))
        except: dto=None
        
        flag=[]
        
        if dfrom is not None and dfrom>100: flag.append('large_origin_snap')
        
        if dto is not None and dto>100: flag.append('large_destination_snap')
        
        if ape is not None and ape>0.15: flag.append('distance_deviation_gt15pct')
        
        p["qa_status"] = "PASS" if not flag else "WATCH"
        p["qa_flags"] = "; ".join(flag)
        
        qa.append(
            {
                "route_id": rid,
                "distKM": dist,
                "observed_km": obs,
                "difference_km": diff,
                "abs_pct_error": ape,
                "dFromKM": dfrom,
                "dToKM": dto,
                "n_chokepoints_final": len(hits),
                "n_chokepoints_exact": len(exact_hits),
                "n_chokepoints_added_by_buffer": len(hits)-len(exact_hits),
                "qa_status": p["qa_status"],
                "qa_flags": p["qa_flags"],
            }
        )
    
    feats.sort(key=lambda ft: route_num(ft.get('properties',{}).get('route_id')))
    
    out_geo = args.out_prefix.with_suffix(".geojson")

    with out_geo.open("w", encoding="utf-8") as f:
        json.dump(
            routes,
            f,
            ensure_ascii=False,
            separators=(",", ":"),
        )
    
    wb=Workbook(); ws=wb.active; ws.title='Routes'
    prop_keys=[]
    for ft in feats:
        for k in ft.get('properties',{}):
            if k not in prop_keys: prop_keys.append(k)
    ws.append(prop_keys)
    for ft in feats: ws.append([ft.get('properties',{}).get(k) for k in prop_keys])
    wr=wb.create_sheet('Route-Chokepoint'); rch=list(rc[0].keys()); wr.append(rch)
    for r in rc: wr.append([r[k] for k in rch])
    wq=wb.create_sheet('QA'); qh=list(qa[0].keys()); wq.append(qh)
    for r in sorted(qa,key=lambda x:route_num(x['route_id'])): wq.append([r[k] for k in qh])
    wm=wb.create_sheet('Method');
    for row in [
        ('Item','Description'),
        ('Transit classification','Intersection between each Eurostat SeaRoute geometry and each of the 28 final chokepoint geometries after applying a uniform metric tolerance buffer.'),
        ('Distance buffers',f'Uniform {BUFFER_KM} km buffer around every chokepoint geometry, constructed in a local azimuthal-equidistant projection and transformed back to CRS84.'),
        ('Exact baseline','The original exact-intersection assignments are retained in chokepoints_exact and n_chokepoints_exact for reproducibility and sensitivity comparison.'),
        ('QA','Flags origin/destination network snap >100 km and route-vs-observed distance deviation >15% for manual review; flags do not automatically alter classifications.'),
    ]: wm.append(row)
    
    fill=PatternFill('solid',fgColor='1F4E78'); font=Font(color='FFFFFF',bold=True)
    for sh in wb.worksheets:
        sh.freeze_panes='A2'
        for c in sh[1]: c.fill=fill; c.font=font; c.alignment=Alignment(horizontal='center')
        for col in range(1,sh.max_column+1):
            vals=[str(sh.cell(r,col).value or '') for r in range(1,min(sh.max_row,200)+1)]
            sh.column_dimensions[get_column_letter(col)].width=min(max(10,max(map(len,vals))+2),45)
        for row in sh.iter_rows():
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    
    out_xlsx = args.out_prefix.parent / "LNG_1037_routes_chokepoint_QA.xlsx"
    wb.save(out_xlsx)
    
    print(f'Wrote {out_geo} and {out_xlsx}; route-chokepoint rows={len(rc)}; flagged routes={sum(bool(x["qa_flags"]) for x in qa)}')

if __name__=='__main__': main()
